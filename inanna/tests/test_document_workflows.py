from __future__ import annotations

import subprocess
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch

import fitz
import openpyxl

from core.document_workflows import (
    DocumentComprehension,
    DocumentDirectReader,
    DocumentRecord,
    DocumentWorkflows,
    DocumentWriter,
    build_document_comprehension,
)
from main import (
    DOCUMENT_TOOL_NAMES,
    build_document_audit_entry,
    build_document_context_lines,
    build_tool_result_text,
    detect_document_tool_action,
    execute_tool_request,
)


class DocumentWorkflowTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.reader = DocumentDirectReader()
        self.writer = DocumentWriter()
        self.desktop = Mock()
        self.desktop.open_app.return_value = Mock(success=True)
        self.workflows = DocumentWorkflows(self.desktop)

    def test_read_text_document_extracts_content(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("alpha beta gamma", encoding="utf-8")

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertEqual(record.format, "txt")
        self.assertEqual(record.word_count, 3)

    def test_read_markdown_document_extracts_title(self) -> None:
        path = self.root / "notes.md"
        path.write_text("# Heading\n\nSecond line", encoding="utf-8")

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertEqual(record.title, "notes")
        self.assertIn("Heading", record.content)

    def test_read_missing_document_returns_error(self) -> None:
        record = self.reader.read(self.root / "missing.pdf")

        self.assertFalse(record.success)
        self.assertIn("File not found", record.error or "")

    def test_read_unsupported_extension_returns_error(self) -> None:
        path = self.root / "blob.bin"
        path.write_bytes(b"123")

        record = self.reader.read(path)

        self.assertFalse(record.success)
        self.assertIn("Unsupported format", record.error or "")

    def test_read_docx_document_extracts_heading(self) -> None:
        docx = __import__("docx")
        path = self.root / "report.docx"
        document = docx.Document()
        document.add_heading("Quarterly Report", level=1)
        document.add_paragraph("Revenue increased steadily.")
        document.save(path)

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertEqual(record.title, "Quarterly Report")
        self.assertIn("Revenue increased", record.content)

    def test_read_pdf_document_extracts_text(self) -> None:
        path = self.root / "brief.pdf"
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 72), "Document faculty proof")
        document.save(path)
        document.close()

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertEqual(record.page_count, 1)
        self.assertIn("Document faculty proof", record.content)

    def test_read_xlsx_document_extracts_sheets_and_cells(self) -> None:
        path = self.root / "sheet.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Alpha"
        sheet["A1"] = "Name"
        sheet["B1"] = "Value"
        sheet["A2"] = "Beta"
        sheet["B2"] = 42
        workbook.save(path)
        workbook.close()

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertEqual(record.sheet_names, ["Alpha"])
        self.assertIn("Beta", record.content)

    def test_read_csv_document_extracts_rows(self) -> None:
        path = self.root / "items.csv"
        path.write_text("name,value\nalpha,1\nbeta,2\n", encoding="utf-8")

        record = self.reader.read(path)

        self.assertTrue(record.success)
        self.assertIn("alpha", record.content)

    def test_build_document_comprehension_collects_summary_and_actions(self) -> None:
        record = DocumentRecord(
            path="report.pdf",
            title="Report",
            format="pdf",
            content="# Intro\n\nThis is a long paragraph that should be summarized for CROWN.\n\n1. First action",
            word_count=2501,
            page_count=3,
        )

        comprehension = build_document_comprehension(record)

        self.assertIn("This is a long paragraph", " ".join(comprehension.summary_lines))
        self.assertIn("1. First action", comprehension.key_points)
        self.assertIn("Ask for a shorter summary", comprehension.suggested_actions)

    def test_document_comprehension_to_crown_context_formats_sections(self) -> None:
        comprehension = DocumentComprehension(
            title="Plan",
            format="docx",
            word_count=120,
            page_count=2,
            summary_lines=["Line one"],
            key_points=["Point one"],
            suggested_actions=["Action one"],
        )

        text = comprehension.to_crown_context()

        self.assertIn("DOCUMENT: Plan (docx)", text)
        self.assertIn("SUMMARY:", text)
        self.assertIn("KEY POINTS:", text)
        self.assertIn("SUGGESTED ACTIONS:", text)

    def test_write_text_document_creates_file(self) -> None:
        path = self.root / "draft.txt"

        result = self.writer.write_text(path, "alpha beta")

        self.assertTrue(result.success)
        self.assertEqual(path.read_text(encoding="utf-8"), "alpha beta")

    def test_write_docx_document_creates_file(self) -> None:
        path = self.root / "draft.docx"

        result = self.writer.write_docx(path, title="Letter", content="Hello world")

        self.assertTrue(result.success)
        self.assertTrue(path.exists())

    def test_export_pdf_uses_libreoffice_cli(self) -> None:
        source = self.root / "draft.docx"
        source.write_text("placeholder", encoding="utf-8")
        output = self.root / "draft.pdf"

        def fake_run(*args, **kwargs):  # type: ignore[no-untyped-def]
            output.write_text("pdf", encoding="utf-8")
            return Mock(returncode=0)

        with patch("core.document_workflows.subprocess.run", side_effect=fake_run) as run_mock:
            result = self.writer.export_pdf_via_libreoffice(source)

        self.assertTrue(result.success)
        self.assertEqual(Path(result.path), output)
        run_mock.assert_called_once()

    def test_export_pdf_timeout_returns_error(self) -> None:
        source = self.root / "draft.docx"
        source.write_text("placeholder", encoding="utf-8")

        with patch(
            "core.document_workflows.subprocess.run",
            side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=30),
        ):
            result = self.writer.export_pdf_via_libreoffice(source)

        self.assertFalse(result.success)
        self.assertIn("timed out", result.error or "")

    def test_workflow_open_in_libreoffice_uses_desktop_faculty(self) -> None:
        source = self.root / "draft.docx"
        source.write_text("placeholder", encoding="utf-8")

        success = self.workflows.open_in_libreoffice(str(source))

        self.assertTrue(success)
        self.desktop.open_app.assert_called_once()

    def test_detect_document_read_action(self) -> None:
        action = detect_document_tool_action("read C:/docs/report.pdf")

        self.assertIsNotNone(action)
        assert action is not None
        self.assertEqual(action["tool"], "doc_read")
        self.assertFalse(action["requires_proposal"])

    def test_detect_document_export_action_requires_proposal(self) -> None:
        action = detect_document_tool_action("export C:/docs/report.docx to pdf")

        self.assertIsNotNone(action)
        assert action is not None
        self.assertEqual(action["tool"], "doc_export_pdf")
        self.assertTrue(action["requires_proposal"])

    def test_execute_tool_request_routes_document_read(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("alpha beta gamma", encoding="utf-8")
        operator = Mock()

        result = execute_tool_request(
            "doc_read",
            {"path": str(path)},
            operator,
            document_workflows=self.workflows,
        )

        self.assertTrue(result.success)
        self.assertEqual(result.tool, "doc_read")
        self.assertIn("alpha beta gamma", result.data["content"])

    def test_build_document_context_lines_include_summary(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("alpha beta gamma\ndelta epsilon zeta", encoding="utf-8")
        result = execute_tool_request(
            "doc_read",
            {"path": str(path)},
            Mock(),
            document_workflows=self.workflows,
        )

        lines = build_document_context_lines(result)

        self.assertTrue(any("content_excerpt:" in line for line in lines))
        self.assertTrue(any("word_count:" in line for line in lines))

    def test_build_document_audit_entry_describes_read(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("alpha beta gamma", encoding="utf-8")
        result = execute_tool_request(
            "doc_read",
            {"path": str(path)},
            Mock(),
            document_workflows=self.workflows,
        )

        audit = build_document_audit_entry(result)

        self.assertIsNotNone(audit)
        assert audit is not None
        self.assertEqual(audit["tool"], "doc_read")
        self.assertIn("words", str(audit["result"]))

    def test_build_tool_result_text_uses_document_formatting(self) -> None:
        path = self.root / "notes.txt"
        path.write_text("alpha beta gamma", encoding="utf-8")
        result = execute_tool_request(
            "doc_read",
            {"path": str(path)},
            Mock(),
            document_workflows=self.workflows,
        )

        text = build_tool_result_text(result)

        self.assertIn("CONTENT", text)
        self.assertIn("DOCUMENT:", text)

    def test_document_tool_names_cover_four_registered_tools(self) -> None:
        self.assertEqual(
            DOCUMENT_TOOL_NAMES,
            {"doc_read", "doc_write", "doc_open", "doc_export_pdf"},
        )


if __name__ == "__main__":
    unittest.main()
